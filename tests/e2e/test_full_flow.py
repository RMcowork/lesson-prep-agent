"""One whole flow through the real code, with the model faked. Slower than the rest; run before a release
(pytest -m e2e), not on every push (practice.md)."""
import threading

import docx
import pytest

import agent.__main__ as app
from agent.telegram_bot import Bot
from agent.settings import Settings
from agent.topics import TopicStore
from tests.helpers import FakeProvider, FakeTelegramAPI, ScriptedChannel, call, docx_call, reply, run_bot_until_done

pytestmark = pytest.mark.e2e
CFG = Settings(max_turns=10)


def lesson_script():
    return [
        reply("", call("ask_user", question="Outline OK?", options=["Approve", "Change something"])),
        reply("", docx_call("lesson-plan", title="خطة الدرس", text="أنواع البيانات", rtl=True)),
        reply("Created the lesson plan."),
        reply("", call("ask_user", question="Save to the knowledge base?", options=["Yes", "No"])),
        reply("", call("write_memory", name="lesson-log.md", content="| 2026 | Data types | Arabic |")),
        reply("Goodbye."),
    ]


def test_terminal_menu_to_files_to_memory(data_dir, monkeypatch):
    for k in ("SMTP_USER", "SMTP_PASSWORD"):
        monkeypatch.delenv(k, raising=False)
    # main menu: Prepare a lesson | wizard: topic 2, 45 min, Arabic, K10, mixed, default outputs, web No, no notes, Start
    # agent: outline Approve | what next: finish | final phase: save? No | main menu: Exit
    ch = ScriptedChannel([0, 1, 1, 0, 0, 0, None, 0, 0, 0, 0, "back", 1, "back"])
    provider = FakeProvider(lesson_script())
    app.main_menu(CFG, ch, provider, threading.Lock(), TopicStore(data_dir=data_dir))

    out = next((data_dir / "output").iterdir())
    assert docx.Document(out / "lesson-plan.docx").paragraphs[0].text == "خطة الدرس"
    assert "w:bidi" in docx.Document(out / "lesson-plan.docx").paragraphs[1]._p.xml          # Arabic lesson is right-to-left
    assert "Data types" in provider.calls[0]["messages"][0]["content"][0]["text"]
    assert "web_search" not in provider.calls[0]["tools"]                                    # web search was off
    assert "Data types" in (data_dir / "memory" / "lesson-log.md").read_text()
    assert ch.sent == [["lesson-plan.docx"]]


def test_telegram_chat_to_files_and_goodbye(data_dir, monkeypatch):
    for k in ("SMTP_USER", "SMTP_PASSWORD"):
        monkeypatch.delenv(k, raising=False)
    # buttons tapped in order: lesson, topic 2, 45 min, Arabic, K10, mixed, outputs (+slides, done), web No, no notes,
    # Start, outline Approve, finish ("b"), final-phase save: No, exit ("b")
    api = FakeTelegramAPI(taps=["o0", "o1", "o1", "o0", "o0", "o0", "t3,d", "o0", "o0", "o0", "o0", "b", "o1", "b"])
    provider = FakeProvider(lesson_script())
    bot = Bot(api, lambda ch: app.main_menu(CFG, ch, provider, threading.Lock(), TopicStore(data_dir=data_dir)),
              allowed_ids=[111], timeout=10)
    api.push_text(999, "/start")                        # a stranger is turned away
    api.push_text(111, "/start")
    assert run_bot_until_done(bot, 111, seconds=20)
    assert api.docs == ["lesson-plan.docx"]
    assert any("Your chat id is 999" in t for t in api.texts_to(999)) and provider.calls
    assert api.texts_to(111)[-1].startswith("Session finished")
