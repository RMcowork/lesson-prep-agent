"""spec.md check 11: Arabic and Hebrew files are right-to-left in Word, Excel and PowerPoint."""
import docx
import openpyxl
import pptx

from agent import docs

BLOCKS = [
    {"type": "heading", "text": "الأهداف", "level": 1},
    {"type": "paragraph", "text": "فهم أنواع البيانات"},
    {"type": "bullets", "items": ["رقمية", "فئوية"]},
    {"type": "numbered", "items": ["أ", "ب"]},
    {"type": "table", "rows": [["النوع", "مثال"], ["رقمي", "العمر"]]},
]


def test_word_rtl_every_paragraph_and_table(tmp_path):
    docs.write_docx(tmp_path / "a.docx", "خطة الدرس", BLOCKS, rtl=True)
    d = docx.Document(tmp_path / "a.docx")
    assert d.paragraphs and all("w:bidi" in p._p.xml for p in d.paragraphs)
    assert all("w:rtl" in r._r.xml for p in d.paragraphs for r in p.runs)
    assert "w:bidiVisual" in d.tables[0]._tbl.xml


def test_word_left_to_right_is_untouched(tmp_path):
    docs.write_docx(tmp_path / "a.docx", "Lesson", BLOCKS, rtl=False)
    d = docx.Document(tmp_path / "a.docx")
    assert not any("w:bidi" in p._p.xml for p in d.paragraphs) and "w:bidiVisual" not in d.tables[0]._tbl.xml


def test_word_table_header_is_bold(tmp_path):
    docs.write_docx(tmp_path / "a.docx", "T", BLOCKS)
    header = docx.Document(tmp_path / "a.docx").tables[0].rows[0].cells[0].paragraphs[0]
    assert header.runs and header.runs[0].bold


def test_excel_rtl_formulas_and_sheet_name_limit(tmp_path):
    docs.write_xlsx(tmp_path / "a.xlsx", [{"name": "ورقة" * 20, "rows": [["العمر", "المدينة"], [15, "حيفا"], ["=SUM(A2:A2)", ""]]}], rtl=True)
    ws = openpyxl.load_workbook(tmp_path / "a.xlsx").active
    assert ws.sheet_view.rightToLeft is True and len(ws.title) <= 31 and ws["A3"].value == "=SUM(A2:A2)"
    docs.write_xlsx(tmp_path / "b.xlsx", [{"name": "S", "rows": [["a"]]}], rtl=False)
    assert not openpyxl.load_workbook(tmp_path / "b.xlsx").active.sheet_view.rightToLeft


def test_powerpoint_rtl_and_notes(tmp_path):
    slides = [{"title": "أنواع البيانات", "bullets": ["رقمية", "فئوية"], "notes": "ملاحظات"}]
    docs.write_pptx(tmp_path / "a.pptx", slides, rtl=True)
    s = pptx.Presentation(tmp_path / "a.pptx").slides[0]
    assert 'rtl="1"' in s.shapes.title._element.xml and 'rtl="1"' in s.placeholders[1]._element.xml
    assert s.notes_slide.notes_text_frame.text == "ملاحظات"
    docs.write_pptx(tmp_path / "b.pptx", [{"title": "T", "bullets": ["x"]}], rtl=False)
    assert 'rtl="1"' not in pptx.Presentation(tmp_path / "b.pptx").slides[0].shapes.title._element.xml
