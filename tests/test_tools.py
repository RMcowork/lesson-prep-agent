"""Tools: web gating (spec check 3), memory safety, versions kept on rewrite (check 4), saving to the knowledge base."""
import pathlib
import socket

import pytest

from agent import memory, tools
from agent.tools import TOOLS, ToolRunner
from tests.helpers import ScriptedChannel


@pytest.fixture
def runner(tmp_path, monkeypatch):
    monkeypatch.setattr(memory, "MEMORY_DIR", tmp_path / "memory")
    return ToolRunner(tmp_path / "out", ScriptedChannel())


# --- web gating -------------------------------------------------------------------------------------
def test_web_off_means_the_model_is_not_offered_web_tools(tmp_path):
    r = ToolRunner(tmp_path, ScriptedChannel(), allow_web=False)
    names = {t["name"] for t in r.tools}
    assert "web_search" not in names and "fetch_url" not in names and "search_knowledge" in names


def test_web_off_refuses_a_web_call_even_if_the_model_makes_one(tmp_path):
    r = ToolRunner(tmp_path, ScriptedChannel(), allow_web=False)
    for name, args in (("web_search", {"query": "x"}), ("fetch_url", {"url": "https://example.com"})):
        with pytest.raises(ValueError, match="turned off"):
            r.run(name, args)


def test_web_on_offers_the_web_tools(tmp_path):
    names = {t["name"] for t in ToolRunner(tmp_path, ScriptedChannel(), allow_web=True).tools}
    assert {"web_search", "fetch_url"} <= names


def test_there_is_no_email_tool_for_the_model():
    """Email is sent by code after the teacher approves, never by the model (architecture.md decision)."""
    assert not [t["name"] for t in TOOLS if "mail" in t["name"] or "send" in t["name"]]


@pytest.mark.parametrize("url", ["http://127.0.0.1:8000", "http://localhost/x", "http://10.0.0.5/", "http://192.168.1.1/",
                                 "http://169.254.169.254/latest", "file:///etc/passwd", "ftp://example.com/x", "not a url"])
def test_fetch_url_refuses_private_and_odd_addresses(runner, url):
    assert runner.run("fetch_url", {"url": url}).startswith("Refused")


def test_fetch_url_marks_content_untrusted_and_truncates(runner, monkeypatch):
    import httpx

    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("93.184.216.34", 0))])

    class Resp:
        text = "<html><script>evil()</script><p>" + "word " * 5000 + "</p></html>"

    monkeypatch.setattr(httpx, "get", lambda *a, **k: Resp())
    out = runner.run("fetch_url", {"url": "https://example.com"})
    assert out.startswith("UNTRUSTED WEB CONTENT") and "evil" not in out and len(out) <= 8100


def test_web_search_falls_back_and_then_degrades_gracefully(runner, monkeypatch):
    import ddgs

    calls = []

    class FakeDDGS:
        def __init__(self, timeout=None): pass
        def text(self, query, max_results, backend):
            calls.append(backend)
            if backend == "duckduckgo":
                raise RuntimeError("blocked")
            return [{"title": "T", "href": "https://x.test", "body": "B"}]

    monkeypatch.setattr(ddgs, "DDGS", FakeDDGS)
    out = runner.run("web_search", {"query": "q"})
    assert calls == ["duckduckgo", "auto"] and out.startswith("UNTRUSTED WEB RESULTS") and "https://x.test" in out

    class Dead(FakeDDGS):
        def text(self, *a, **k): raise RuntimeError("down")

    monkeypatch.setattr(ddgs, "DDGS", Dead)
    assert "No web results" in runner.run("web_search", {"query": "q"})


# --- memory (the safe operations are tested in tests/test_memory.py) --------------------------------
def test_the_model_can_add_to_memory_but_not_replace_a_file(runner):
    names = {t["name"] for t in runner.tools}
    assert "add_memory" in names and "read_memory" in names
    assert not {n for n in names if "write_memory" in n or "rewrite" in n or "replace" in n}
    with pytest.raises(ValueError, match="Unknown tool"):
        runner.run("write_memory", {"name": "lesson-log.md", "content": "everything"})


def test_memory_round_trip_through_the_tools_and_only_three_files(runner):
    assert runner.run("read_memory", {"name": "lesson-log.md"}) == "(empty)"
    runner.run("add_memory", {"name": "teacher-preferences.md", "entry": "prefers short warm-ups"})
    assert "- prefers short warm-ups" in runner.run("read_memory", {"name": "teacher-preferences.md"})
    for bad in ("secrets.md", "INDEX.md"):
        with pytest.raises(ValueError):
            runner.run("add_memory", {"name": bad, "entry": "x"})


# --- files and versions (spec check 4) --------------------------------------------------------------
def docx(runner, text):
    return runner.run("create_docx", {"filename": "plan", "title": "T", "blocks": [{"type": "paragraph", "text": text}]})


def test_rewriting_a_file_keeps_each_earlier_version(runner):
    import docx as d

    docx(runner, "v1"); docx(runner, "v2"); docx(runner, "v3")
    out = runner.out_dir
    assert d.Document(out / "plan.docx").paragraphs[1].text == "v3"
    backups = sorted(p.name for p in (out / "previous-versions").iterdir())
    assert backups == ["plan-v1.docx", "plan-v2.docx"]
    assert d.Document(out / "previous-versions" / "plan-v1.docx").paragraphs[1].text == "v1"


def test_filenames_cannot_escape_the_job_folder(runner, tmp_path):
    runner.run("create_docx", {"filename": "../../evil", "title": "T", "blocks": []})
    assert (runner.out_dir / "evil.docx").exists() and not (tmp_path.parent / "evil.docx").exists()


def test_xlsx_and_pptx_are_created(runner):
    runner.run("create_xlsx", {"filename": "a", "sheets": [{"name": "S", "rows": [["h"], [1]]}]})
    runner.run("create_pptx", {"filename": "b", "slides": [{"title": "Hi", "bullets": ["x"]}]})
    assert (runner.out_dir / "a.xlsx").exists() and (runner.out_dir / "b.pptx").exists()


# --- saving to the knowledge base -------------------------------------------------------------------
@pytest.fixture
def kb(tmp_path, monkeypatch):
    from rag import core

    monkeypatch.setattr(core, "KNOWLEDGE", tmp_path / "knowledge")
    monkeypatch.setattr(core, "reindex", lambda: "indexed")
    return core


def test_add_to_knowledge_copies_approved_files_under_generated(runner, kb):
    docx(runner, "final")
    msg = runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    saved = kb.KNOWLEDGE / "generated" / runner.out_dir.name / "plan.docx"
    assert saved.exists() and "indexed" in msg


def test_add_to_knowledge_refuses_files_outside_the_job_folder(runner, kb, tmp_path):
    (tmp_path / "secret.txt").write_text("private")
    with pytest.raises(ValueError):
        runner.run("add_to_knowledge", {"filenames": ["../secret.txt"]})
    assert not any((kb.KNOWLEDGE / "generated").rglob("secret.txt")) if (kb.KNOWLEDGE / "generated").exists() else True


def test_add_to_knowledge_never_overwrites(runner, kb):
    """Ticket #40: a second save with the same file name used to replace the first one, with no backup."""
    import docx as d

    saved = kb.KNOWLEDGE / "generated" / runner.out_dir.name
    docx(runner, "first version")
    runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    docx(runner, "second version")                                           # the teacher asked for a change, then saves again
    msg = runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    assert "plan.docx as plan-v2.docx" in msg
    assert d.Document(saved / "plan.docx").paragraphs[1].text == "first version"        # the first save is untouched
    assert d.Document(saved / "plan-v2.docx").paragraphs[1].text == "second version"
    docx(runner, "third version")
    assert "plan.docx as plan-v3.docx" in runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    assert sorted(p.name for p in saved.iterdir()) == ["plan-v2.docx", "plan-v3.docx", "plan.docx"]


def test_saving_identical_content_again_does_nothing_and_says_so(runner, kb):
    saved = kb.KNOWLEDGE / "generated" / runner.out_dir.name
    docx(runner, "same")
    runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    msg = runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    assert "Already saved, identical: plan.docx" in msg and "Saved plan" not in msg
    assert sorted(p.name for p in saved.iterdir()) == ["plan.docx"]


def test_a_save_of_several_files_reports_each_one(runner, kb):
    runner.run("create_xlsx", {"filename": "task", "sheets": [{"name": "S", "rows": [["a"]]}]})
    docx(runner, "v1")
    runner.run("add_to_knowledge", {"filenames": ["plan.docx"]})
    docx(runner, "v2")
    msg = runner.run("add_to_knowledge", {"filenames": ["plan.docx", "task.xlsx"]})
    assert "Saved plan.docx as plan-v2.docx, task.xlsx" in msg


# --- asking the teacher -----------------------------------------------------------------------------
def test_ask_user_with_options_returns_the_chosen_option(tmp_path):
    r = ToolRunner(tmp_path, ScriptedChannel([1]))
    assert r.run("ask_user", {"question": "ok?", "options": ["Approve", "Change"]}) == "Change"


def test_ask_user_typed_answer_and_free_text(tmp_path):
    r = ToolRunner(tmp_path, ScriptedChannel([("other", "shorter please"), "what to change"]))
    assert r.run("ask_user", {"question": "ok?", "options": ["Approve"]}) == "shorter please"
    assert r.run("ask_user", {"question": "What should change?"}) == "what to change"


def test_unknown_tool_is_an_error(tmp_path):
    with pytest.raises(ValueError):
        ToolRunner(tmp_path, ScriptedChannel()).run("rm_rf", {})
