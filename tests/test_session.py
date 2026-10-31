"""spec.md check 4: a change note rewrites the file under the same name, keeps the old version, and the
final step (save to knowledge base, memory) runs only after the teacher finishes."""
import json
import threading

import docx
import pytest

import agent.__main__ as app
from agent.settings import Settings
from tests.helpers import FakeProvider, ScriptedChannel, call, docx_call, reply

CFG = Settings(max_turns=10)


@pytest.fixture(autouse=True)
def no_email(monkeypatch):
    for k in ("SMTP_USER", "SMTP_PASSWORD", "EMAIL_TO"):
        monkeypatch.delenv(k, raising=False)


def session_script():
    return [
        reply("", docx_call("plan", text="version 1")), reply("Created the plan."),                      # first turn
        reply("", docx_call("plan", text="version 2 shorter warm-up")), reply("Changed the warm-up."),   # revision
        reply("", call("ask_user", question="Save to the knowledge base?", options=["Yes", "No"])),      # final phase
        reply("", call("add_memory", name="teacher-preferences.md", entry="prefers short warm-ups")),
        reply("Goodbye."),
    ]


def run(data_dir, ch, provider=None, **kw):
    provider = provider or FakeProvider(session_script())
    app.run_session(provider, CFG, "Prepare a lesson. Topic: X", "revision test", ch, **kw)
    return provider


def test_revision_flow_end_to_end(data_dir):
    ch = ScriptedChannel([0, "make the warm-up shorter", "back", 1])      # change request, note, finish, KB: No
    p = run(data_dir, ch)

    out = next((data_dir / "output").iterdir())
    assert docx.Document(out / "plan.docx").paragraphs[1].text.startswith("version 2")
    assert docx.Document(out / "previous-versions" / "plan-v1.docx").paragraphs[1].text == "version 1"
    assert ch.sent == [["plan.docx"], ["plan.docx"]]                       # delivered, then delivered again after the change

    note_call = p.calls[2]["messages"]                                      # 3rd model call = first of the revision
    assert note_call[-1]["content"][0]["text"] == "make the warm-up shorter"
    assert "short warm-ups" in (data_dir / "memory" / "teacher-preferences.md").read_text()


def test_the_final_step_starts_only_after_the_teacher_finishes(data_dir):
    ch = ScriptedChannel([0, "a note", "back", 1])
    p = run(data_dir, ch)
    texts = lambda c: [b.get("text", "") for m in c["messages"] if m["role"] == "user" for b in m["content"] if b["type"] == "text"]
    assert not any("satisfied" in t for call_ in p.calls[:4] for t in texts(call_))    # never mentioned before finishing
    assert any("satisfied" in t for t in texts(p.calls[4]))                              # first final-phase call


def test_finishing_straight_away_still_runs_the_final_step(data_dir):
    script = [reply("", docx_call()), reply("Created."),
              reply("", call("ask_user", question="Save?", options=["Yes", "No"])), reply("Bye.")]
    ch = ScriptedChannel(["back", 1])
    run(data_dir, ch, FakeProvider(script))
    assert ch.sent == [["plan.docx"]]


def test_token_totals_are_written_next_to_the_log(data_dir):
    run(data_dir, ScriptedChannel([0, "note", "back", 1]))
    usage = json.loads(next((data_dir / "logs").glob("*.usage.json")).read_text())
    assert usage == {"calls": 7, "input_tokens": 700, "output_tokens": 140}


def test_no_files_created_is_said_plainly_and_the_session_continues(data_dir):
    script = [reply("I need a topic first."), reply("Goodbye.")]
    ch = ScriptedChannel(["back"])
    run(data_dir, ch, FakeProvider(script))
    assert any("No files were created or changed." in t for t in ch.said)


def test_the_lock_is_released_after_each_agent_turn(data_dir):
    lock = threading.Lock()
    ch = ScriptedChannel([0, "note", "back", 1])
    run(data_dir, ch, lock=lock)
    assert not lock.locked()


def test_a_waiting_job_tells_the_teacher_and_then_runs(data_dir):
    lock = threading.Lock()
    lock.acquire()
    threading.Timer(0.3, lock.release).start()
    ch = ScriptedChannel([0, "note", "back", 1])
    run(data_dir, ch, lock=lock)
    assert any("Another job is running" in t for t in ch.said)


def test_web_off_is_enforced_for_the_whole_session(data_dir):
    p = FakeProvider([reply("", docx_call()), reply("done"),
                      reply("", call("ask_user", question="Save?", options=["Yes", "No"])), reply("bye")])
    run(data_dir, ScriptedChannel(["back", 1]), p, allow_web=False)
    assert all("web_search" not in c["tools"] and "fetch_url" not in c["tools"] for c in p.calls)


def one_session(data_dir, label):
    script = [reply("", docx_call("plan", text=label)), reply("done"),
              reply("", call("ask_user", question="Save?", options=["Yes", "No"])), reply("bye")]
    run(data_dir, ScriptedChannel(["back", 1]), FakeProvider(script))


def test_two_sessions_on_the_same_topic_get_their_own_folders_and_logs(data_dir):
    """Bug #29: the second session used to overwrite the first one's log and share its folder."""
    one_session(data_dir, "FIRST session content")
    one_session(data_dir, "SECOND session content")
    one_session(data_dir, "THIRD session content")
    folders = sorted(p.name for p in (data_dir / "output").iterdir())
    assert len(folders) == 3 and folders[0].endswith("revision-test") and folders[1].endswith("revision-test-2") and folders[2].endswith("revision-test-3")
    logs = sorted(p.name for p in (data_dir / "logs").glob("*.json") if "usage" not in p.name)
    assert set(logs) == {f + ".json" for f in folders} and len(logs) == 3            # exactly one log per session
    for folder, label in zip(folders, ("FIRST", "SECOND", "THIRD")):
        assert f"{label} session content" in (data_dir / "logs" / f"{folder}.json").read_text()   # no log was overwritten
        import docx as d
        assert d.Document(data_dir / "output" / folder / "plan.docx").paragraphs[1].text == f"{label} session content"
        assert not (data_dir / "output" / folder / "previous-versions").exists()    # each session started clean
    assert len(list((data_dir / "logs").glob("*.usage.json"))) == 3


def test_a_revision_inside_one_session_stays_in_that_sessions_folder(data_dir):
    run(data_dir, ScriptedChannel([0, "make it shorter", "back", 1]))
    assert len(list((data_dir / "output").iterdir())) == 1                           # one session, one folder, however many revisions


def test_a_leftover_log_alone_is_enough_to_make_a_name_unavailable(data_dir):
    import agent.__main__ as app

    (data_dir / "output" / "2026-10-31-x").mkdir(parents=True)
    (data_dir / "logs").mkdir()
    (data_dir / "logs" / "2026-10-31-x-2.json").write_text("[]")              # a log whose folder was deleted
    assert app.unique_folder("2026-10-31-x") == "2026-10-31-x-3"
    assert app.unique_folder("2026-10-31-y") == "2026-10-31-y"                # an unused name is kept as it is


def test_numbered_folder_names_still_sort_by_date():
    assert sorted(["2026-10-31-x-2", "2026-10-30-x", "2026-10-31-x"]) == ["2026-10-30-x", "2026-10-31-x", "2026-10-31-x-2"]
