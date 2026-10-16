"""Local RAG core: index and search the teacher's knowledge/ folder.

Used by both the MCP server (rag/server.py) and the standalone agent.
"""
import hashlib
import json
import os
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = Path(os.environ.get("AGENT_DATA_DIR", ROOT)).resolve()  # mount a volume here when deployed
KNOWLEDGE = DATA / "knowledge"
INDEX_DIR = DATA / "rag-index"
MANIFEST = INDEX_DIR / "manifest.json"
MODEL_NAME = "intfloat/multilingual-e5-large"
CHUNK_WORDS = 250
OVERLAP_WORDS = 40
SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv"}

_model = None
_collection = None


def _get_model():
    global _model
    if _model is None:
        from fastembed import TextEmbedding

        _model = TextEmbedding(
            model_name=MODEL_NAME, cache_dir=str(INDEX_DIR / "models")
        )
    return _model


def _get_collection():
    global _collection
    if _collection is None:
        import chromadb

        INDEX_DIR.mkdir(parents=True, exist_ok=True)
        client = chromadb.PersistentClient(path=str(INDEX_DIR / "chroma"))
        _collection = client.get_or_create_collection(
            "knowledge", metadata={"hnsw:space": "cosine"}
        )
    return _collection


def _embed(texts, prefix):
    # e5 models expect "query: " / "passage: " prefixes; fastembed returns normalized vectors
    return [v.tolist() for v in _get_model().embed([f"{prefix}: {t}" for t in texts])]


def extract(path: Path):
    """Return a list of (location, text) pairs for a file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return [
            (f"page {i}", page.extract_text() or "")
            for i, page in enumerate(PdfReader(str(path)).pages, 1)
        ]
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [("document", "\n".join(parts))]
    if suffix == ".pptx":
        from pptx import Presentation

        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame
            ]
            out.append((f"slide {i}", "\n".join(texts)))
        return out
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [
                " | ".join("" if v is None else str(v) for v in row)
                for row in ws.iter_rows(values_only=True)
            ]
            out.append((f"sheet {ws.title}", "\n".join(rows)))
        return out
    return [("file", path.read_text(encoding="utf-8", errors="ignore"))]


def chunk(text: str):
    words = text.split()
    step = CHUNK_WORDS - OVERLAP_WORDS
    for start in range(0, max(len(words), 1), step):
        piece = " ".join(words[start : start + CHUNK_WORDS])
        if piece.strip():
            yield piece
        if start + CHUNK_WORDS >= len(words):
            break


def _file_hash(path: Path) -> str:
    return hashlib.sha1(path.read_bytes()).hexdigest()


def _load_manifest() -> dict:
    if MANIFEST.exists():
        return json.loads(MANIFEST.read_text())
    return {}


def reindex() -> str:
    """Index new/changed files in knowledge/ and drop deleted ones."""
    KNOWLEDGE.mkdir(exist_ok=True)
    INDEX_DIR.mkdir(parents=True, exist_ok=True)
    col = _get_collection()
    manifest = _load_manifest()
    files = {
        str(p.relative_to(KNOWLEDGE)): p
        for p in KNOWLEDGE.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith(".")
    }
    added = updated = removed = 0
    errors = []

    for rel in list(manifest):
        if rel not in files:
            col.delete(where={"file": rel})
            del manifest[rel]
            removed += 1

    for rel, path in files.items():
        digest = _file_hash(path)
        if manifest.get(rel) == digest:
            continue
        try:
            col.delete(where={"file": rel})
            ids, docs, metas = [], [], []
            topic = rel.split("/")[0] if "/" in rel else ""
            source = "generated" if topic == "generated" else "teacher"
            for location, text in extract(path):
                for n, piece in enumerate(chunk(text)):
                    ids.append(f"{rel}::{location}::{n}")
                    docs.append(piece)
                    metas.append(
                        {"file": rel, "location": location, "topic": topic, "source": source}
                    )
            if docs:
                col.add(
                    ids=ids,
                    documents=docs,
                    metadatas=metas,
                    embeddings=_embed(docs, "passage"),
                )
            updated += 1 if rel in manifest else 0
            added += 0 if rel in manifest else 1
            manifest[rel] = digest
        except Exception as exc:  # keep going; report at the end
            errors.append(f"{rel}: {exc}")

    MANIFEST.write_text(json.dumps(manifest, indent=2))
    msg = f"added {added}, updated {updated}, removed {removed}; {col.count()} chunks total"
    if errors:
        msg += "\nerrors:\n" + "\n".join(errors)
    return msg


def search_knowledge(query: str, topic: str = "", k: int = 5, source: str = "") -> str:
    """Search the teacher's course material. Returns chunks with file and page/slide.

    topic optionally limits results to a subfolder of knowledge/.
    """
    col = _get_collection()
    if col.count() == 0:
        return "Knowledge base is empty. Add files to knowledge/ and run reindex."
    filters = []
    if topic:
        filters.append({"topic": topic})
    if source in ("teacher", "generated"):
        filters.append({"source": source})
    where = None if not filters else filters[0] if len(filters) == 1 else {"$and": filters}
    res = col.query(
        query_embeddings=_embed([query], "query"),
        n_results=min(k, col.count()),
        where=where,
    )
    if not res["documents"][0]:
        return "No relevant material found."
    out = []
    for doc, meta, dist in zip(
        res["documents"][0], res["metadatas"][0], res["distances"][0]
    ):
        score = round(1 - dist, 3)
        tag = "GENERATED earlier by this agent" if meta.get("source") == "generated" else "teacher material"
        out.append(f"[{meta['file']} | {meta['location']} | {tag} | score {score}]\n{doc}")
    return "\n\n---\n\n".join(out)

