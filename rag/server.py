"""Local RAG MCP server for the lesson-prep agent.

Tools:
  search_knowledge(query, topic="", k=5)  -> relevant chunks from knowledge/
  reindex()                               -> (re)index new, changed and deleted files
"""
import hashlib
import json
from pathlib import Path

from mcp.server.fastmcp import FastMCP

ROOT = Path(__file__).resolve().parent.parent
KNOWLEDGE = ROOT / "knowledge"
INDEX_DIR = Path(__file__).resolve().parent / "index"
MANIFEST = INDEX_DIR / "manifest.json"
MODEL_NAME = "intfloat/multilingual-e5-large"
CHUNK_WORDS = 250
OVERLAP_WORDS = 40
SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv"}

mcp = FastMCP("rag")
_model = None
_collection = None


def _get_model():
    global _model
    if _model is None:
        from fastembed import TextEmbedding

        _model = TextEmbedding(model_name=MODEL_NAME)
    return _model


def _get_collection():
    global _collection
    if _collection is None:
        import chromadb

        INDEX_DIR.mkdir(parents=True, exist_ok=True)
        client = chromadb.PersistentClient(path=str(INDEX_DIR / "chroma"))
        _collection = client.get_or_create_collection(
            "knowledge", metadata={"hnsw:space": "cosine"}
        )
    return _collection


def _embed(texts, prefix):
    # e5 models expect "query: " / "passage: " prefixes; fastembed returns normalized vectors
    return [v.tolist() for v in _get_model().embed([f"{prefix}: {t}" for t in texts])]


def extract(path: Path):
    """Return a list of (location, text) pairs for a file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return [
            (f"page {i}", page.extract_text() or "")
            for i, page in enumerate(PdfReader(str(path)).pages, 1)
        ]
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [("document", "\n".join(parts))]
    if suffix == ".pptx":
        from pptx import Presentation

        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame
            ]
            out.append((f"slide {i}", "\n".join(texts)))
        return out
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [
                " | ".join("" if v is None else str(v) for v in row)
                for row in ws.iter_rows(values_only=True)
            ]
            out.append((f"sheet {ws.title}", "\n".join(rows)))
        return out
    return [("file", path.read_text(encoding="utf-8", errors="ignore"))]


def chunk(text: str):
    words = text.split()
    step = CHUNK_WORDS - OVERLAP_WORDS
    for start in range(0, max(len(words), 1), step):
        piece = " ".join(words[start : start + CHUNK_WORDS])
        if piece.strip():
            yield piece
        if start + CHUNK_WORDS >= len(words):
            break


def _file_hash(path: Path) -> str:
    return hashlib.sha1(path.read_bytes()).hexdigest()


def _load_manifest() -> dict:
    if MANIFEST.exists():
        return json.loads(MANIFEST.read_text())
    return {}


@mcp.tool()
def reindex() -> str:
    """Index new/changed files in knowledge/ and drop deleted ones."""
    KNOWLEDGE.mkdir(exist_ok=True)
    INDEX_DIR.mkdir(parents=True, exist_ok=True)
    col = _get_collection()
    manifest = _load_manifest()
    files = {
        str(p.relative_to(KNOWLEDGE)): p
        for p in KNOWLEDGE.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith(".")
    }
    added = updated = removed = 0
    errors = []

    for rel in list(manifest):
        if rel not in files:
            col.delete(where={"file": rel})
            del manifest[rel]
            removed += 1

    for rel, path in files.items():
        digest = _file_hash(path)
        if manifest.get(rel) == digest:
            continue
        try:
            col.delete(where={"file": rel})
            ids, docs, metas = [], [], []
            topic = rel.split("/")[0] if "/" in rel else ""
            for location, text in extract(path):
                for n, piece in enumerate(chunk(text)):
                    ids.append(f"{rel}::{location}::{n}")
                    docs.append(piece)
                    metas.append({"file": rel, "location": location, "topic": topic})
            if docs:
                col.add(
                    ids=ids,
                    documents=docs,
                    metadatas=metas,
                    embeddings=_embed(docs, "passage"),
                )
            updated += 1 if rel in manifest else 0
            added += 0 if rel in manifest else 1
            manifest[rel] = digest
        except Exception as exc:  # keep going; report at the end
            errors.append(f"{rel}: {exc}")

    MANIFEST.write_text(json.dumps(manifest, indent=2))
    msg = f"added {added}, updated {updated}, removed {removed}; {col.count()} chunks total"
    if errors:
        msg += "\nerrors:\n" + "\n".join(errors)
    return msg


@mcp.tool()
def search_knowledge(query: str, topic: str = "", k: int = 5) -> str:
    """Search the teacher's course material. Returns chunks with file and page/slide.

    topic optionally limits results to a subfolder of knowledge/.
    """
    col = _get_collection()
    if col.count() == 0:
        return "Knowledge base is empty. Add files to knowledge/ and run reindex."
    where = {"topic": topic} if topic else None
    res = col.query(
        query_embeddings=_embed([query], "query"),
        n_results=min(k, col.count()),
        where=where,
    )
    if not res["documents"][0]:
        return "No relevant material found."
    out = []
    for doc, meta, dist in zip(
        res["documents"][0], res["metadatas"][0], res["distances"][0]
    ):
        score = round(1 - dist, 3)
        out.append(f"[{meta['file']} | {meta['location']} | score {score}]\n{doc}")
    return "\n\n---\n\n".join(out)


if __name__ == "__main__":
    mcp.run()
