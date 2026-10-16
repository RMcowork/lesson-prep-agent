"""File generators used by the agent's tools: .docx, .xlsx, .pptx (with right-to-left support)."""
from pathlib import Path


def _rtl_paragraph(paragraph):
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rPr.append(OxmlElement("w:rtl"))
        fonts = OxmlElement("w:rFonts")
        fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs", "Arial")
        rPr.append(fonts)


def write_docx(path: Path, title: str, blocks: list, rtl: bool = False):
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document()
    paragraphs = []
    paragraphs.append(doc.add_heading(title, level=0))
    for block in blocks:
        kind = block.get("type", "paragraph")
        if kind == "heading":
            paragraphs.append(doc.add_heading(block["text"], level=int(block.get("level", 1))))
        elif kind == "paragraph":
            paragraphs.append(doc.add_paragraph(block["text"]))
        elif kind in ("bullets", "numbered"):
            style = "List Bullet" if kind == "bullets" else "List Number"
            for item in block["items"]:
                paragraphs.append(doc.add_paragraph(str(item), style=style))
        elif kind == "table":
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    cell = table.cell(i, j)
                    cell.text = str(value)
                    for p in cell.paragraphs:
                        if i == 0 and block.get("header", True):
                            for run in p.runs:
                                run.bold = True
                        paragraphs.append(p)
            if rtl:
                tblPr = table._tbl.tblPr
                tblPr.append(OxmlElement("w:bidiVisual"))
    if rtl:
        for p in paragraphs:
            _rtl_paragraph(p)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def write_xlsx(path: Path, sheets: list, rtl: bool = False):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(str(sheet["name"])[:31])
        for row in sheet["rows"]:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        if rtl:
            ws.sheet_view.rightToLeft = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def write_pptx(path: Path, slides: list, rtl: bool = False):
    from pptx import Presentation

    prs = Presentation()
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        body = slide.placeholders[1].text_frame
        bullets = spec.get("bullets", [])
        for i, text in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(text)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        if rtl:
            for frame in (slide.shapes.title.text_frame, body):
                for p in frame.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    pPr.set("rtl", "1")
                    pPr.set("algn", "r")
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
